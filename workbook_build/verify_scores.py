
import openpyxl
wb=openpyxl.load_workbook("Mobile App Accessibility Evaluation.xlsx")
sc={'Pass':1,'Partial':.5,'Fail':0}
def band(x): return "Level AA-aligned" if x>=90 else "Substantially accessible" if x>=75 else "Partially accessible - significant barriers" if x>=50 else "Largely inaccessible" if x>=25 else "Inaccessible / excludes blind users"
for n in ['BK Mobile','Equity Mobile RW','Fidelity NG','GTWorld NG','AccessMore NG']:
    ws=wb[n]; rows=[(ws.cell(r,3).value,ws.cell(r,5).value,ws.cell(r,6).value) for r in range(5,40)]
    def pct(lv):
        t=[(w,sc[s]) for l,w,s in rows if s in sc and l in lv]; return 100*sum(w*v for w,v in t)/sum(w for w,v in t)
    o=pct({'A','AA','AAA'}); fails=sum(s=='Fail' for _,_,s in rows); nt=sum(s in('Not Tested','N/A') for _,_,s in rows)
    t=wb[n+' - Transactions']; tr=[(t.cell(r,4).value,t.cell(r,5).value) for r in range(5,17)]
    tt=[(w,sc[s]) for w,s in tr if s in sc]; ts=100*sum(w*v for w,v in tt)/sum(w for w,v in tt) if tt else 0
    tnt=sum(s=='Not Tested' for w,s in tr)
    print(f"{n:18} WCAG {o:5.1f} {band(o):45} A {pct({'A'}):5.1f} AA {pct({'A','AA'}):5.1f} | Txn {ts:5.1f} | Fail {fails} | manual {nt}+{tnt} | scored {35-nt}")
d=wb['Comparison Dashboard']
for r in range(4,12): print([str(c.value)[:40] for c in d[r]])
